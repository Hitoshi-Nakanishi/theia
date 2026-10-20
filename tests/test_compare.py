import csv
import json
import os
import tempfile
import time
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from dataclasses import asdict

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook as ExcelWorkbook
from PySide6.QtCore import QItemSelectionModel, QSettings, Qt
from PySide6.QtWidgets import QApplication

from theia.compare import (
    ComparisonDialog, ComparisonModel, ComparisonProxy, ComparisonRow,
    ComparisonWorker, Metric, csv_value, export_csv,
    cell_locator, load_preset,
)
from theia.workbook import Workbook


APP = QApplication.instance() or QApplication([])


def make_workbook(path, value=10, label="Revenue", header="FY1", sheet="Model"):
    wb = ExcelWorkbook()
    ws = wb.active
    ws.title = sheet
    ws["C3"] = header
    ws["B4"] = label
    ws["C4"] = value
    ws["B5"] = "Margin"
    ws["C5"] = 0.125
    wb.save(path)
    wb.close()


def wait_until(predicate, timeout=15):
    deadline = time.monotonic() + timeout
    while not predicate():
        APP.processEvents()
        if time.monotonic() > deadline:
            raise AssertionError("Timed out waiting for background read")
        time.sleep(0.005)
    APP.processEvents()


class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.reference = self.root / "Company_000.xlsx"
        make_workbook(self.reference)
        self.metric = Metric("Revenue", 3, 2)

    def tearDown(self):
        APP.processEvents()
        self.temp.cleanup()

    def extract(self, paths, metrics=None, check=True):
        worker = ComparisonWorker([(p.stem if isinstance(p, Path) else "Archive", p) for p in paths],
                                  self.reference, "Model", metrics or [self.metric], check)
        rows, errors = [], []
        worker.row_ready.connect(rows.append)
        worker.failed.connect(errors.append)
        worker.run()
        self.assertEqual(errors, [])
        return rows

    def test_coordinates_include_empty_leading_rows_and_columns(self):
        with Workbook(self.reference) as wb:
            rows = wb.load_sheet("Model").rows
            self.assertEqual(rows[3][2], 10)
            self.assertEqual(rows[2][2], "FY1")
        self.assertEqual(self.extract([self.reference])[0].values, [10])

    def test_missing_mismatched_zero_and_formula_values(self):
        paths = []
        for name, kwargs in [
            ("zero", {"value": 0}), ("negative", {"value": -2}),
            ("mismatch", {"label": "Other"}), ("period", {"header": "FY2"}),
            ("blank", {"value": None}), ("formula", {"value": "=2+3"}),
            ("missing", {"sheet": "Other"}),
        ]:
            path = self.root / f"{name}.xlsx"
            make_workbook(path, **kwargs)
            paths.append(path)
        results = {r.label: r for r in self.extract(paths)}
        self.assertEqual(results["zero"].values, [0])
        self.assertEqual(results["negative"].values, [-2])
        for name in ("mismatch", "period"):
            self.assertIsNone(results[name].values[0])
            self.assertIn("mismatch", results[name].status)
        for name in ("blank", "formula"):
            self.assertIsNone(results[name].values[0])
            self.assertIn("no cached value", results[name].status)
        self.assertEqual(results["missing"].status, "Missing sheet")
        self.assertEqual(self.extract([paths[2]], check=False)[0].values, [10])

    def test_bad_file_does_not_stop_batch_and_zip_supported(self):
        broken = self.root / "broken.xlsx"
        broken.write_text("invalid workbook", encoding="utf-8")
        archive = self.root / "models.zip"
        with zipfile.ZipFile(archive, "w") as z:
            z.write(self.reference, "models/Company_000.xlsx")
        results = self.extract([broken, (archive, "models/Company_000.xlsx"), self.reference])
        self.assertIn("Read error", results[0].status)
        self.assertEqual(results[1].values, [10])
        self.assertEqual(results[2].values, [10])

    def test_saved_formula_cache_is_used(self):
        path = self.root / "cached.xlsx"
        make_workbook(path, value="=2+3")
        data = path.read_bytes()
        with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(path, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    xml = ElementTree.fromstring(payload)
                    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    xml.find(".//s:c[@r='C4']/s:v", ns).text = "5"
                    payload = ElementTree.tostring(xml)
                target.writestr(info, payload)
        self.assertEqual(self.extract([path])[0].values, [5])

    def test_workbook_evicts_and_reloads_sheets(self):
        path = self.root / "many_sheets.xlsx"
        wb = ExcelWorkbook()
        for i in range(5):
            wb.create_sheet(f"Sheet{i}")["A1"] = i
        wb.save(path)
        wb.close()
        with Workbook(path) as reader:
            for i in range(5):
                self.assertEqual(reader.load_sheet(f"Sheet{i}").rows[0][0], i)
            self.assertEqual(sum(s.loaded for s in reader.sheets.values()), 1)
            self.assertEqual(reader.load_sheet("Sheet0").rows[0][0], 0)

    def test_match_labels_handles_different_history_lengths_and_row_positions(self):
        def layout(path, count, offset, unit="units", duplicate=False):
            wb = ExcelWorkbook()
            ws = wb.active
            ws.title = "Model"
            ws.cell(3 + offset, 1, "Period")
            ws.cell(4 + offset, 1, "Kind")
            for col in range(3, 3 + count):
                ws.cell(3 + offset, col, f"History {col}")
                ws.cell(4 + offset, col, "Actual")
            for n in range(2):
                col = 3 + count + n
                ws.cell(3 + offset, col, f"Future {n + 1}")
                ws.cell(4 + offset, col, "Forecast")
                ws.cell(6 + offset, col, 100 + n)
            ws.cell(6 + offset, 1, "Revenue")
            ws.cell(6 + offset, 2, unit)
            if duplicate:
                ws.cell(10 + offset, 1, "Revenue")
            wb.save(path)
            wb.close()

        reference = self.root / "reference.xlsx"
        shifted = self.root / "shifted.xlsx"
        mismatch = self.root / "unit_mismatch.xlsx"
        duplicate = self.root / "duplicate.xlsx"
        layout(reference, 2, 0)
        layout(shifted, 5, 3)
        layout(mismatch, 4, 0, unit="other units")
        layout(duplicate, 3, 0, duplicate=True)
        metrics = [Metric("First forecast", 5, 4), Metric("Second forecast", 5, 5)]
        worker = ComparisonWorker([(p.stem, p) for p in (shifted, mismatch, duplicate)],
                                  reference, "Model", metrics, True, match_labels=True)
        results, failures = [], []
        worker.row_ready.connect(results.append)
        worker.failed.connect(failures.append)
        worker.run()
        self.assertEqual(failures, [])
        self.assertEqual(results[0].values, [100, 101])
        self.assertEqual(results[0].coordinates, [(8, 7), (8, 8)])
        self.assertIn("Future 1", results[0].locations[0])
        self.assertEqual(results[1].values, [None, None])
        self.assertIn("unit", results[1].status)
        self.assertIn("ambiguous", results[2].status)

        with Workbook(reference) as wb:
            rows = wb.load_sheet("Model").rows
            saved = [Metric(m.title, m.row, m.col, cell_locator(rows, m)) for m in metrics]
        preset = self.root / "comparison.json"
        preset.write_text(json.dumps({"version": 1, "sheet": "Model", "match_labels": True,
                                      "metrics": [asdict(m) for m in saved]}), encoding="utf-8")
        sheet, loaded, match = load_preset(preset)
        results.clear()
        worker = ComparisonWorker([("Shifted", shifted)], self.root / "absent_reference.xlsx",
                                  sheet, loaded, True, match_labels=match)
        worker.row_ready.connect(results.append)
        worker.run()
        self.assertEqual(results[0].values, [100, 101])
        preset.write_text('{"version":1,"sheet":"Model","metrics":[]}', encoding="utf-8")
        with self.assertRaises(ValueError):
            load_preset(preset)

    def test_sort_filter_and_csv_preserve_raw_numbers(self):
        model = ComparisonModel([self.metric], sheet_name="Model")
        for label, value in [("A", 100), ("B", 9), ("C", -2), ("D", None)]:
            model.append(ComparisonRow(label, self.reference, [value], "OK"))
        proxy = ComparisonProxy()
        proxy.setSourceModel(model)
        proxy.setFilterKeyColumn(-1)
        proxy.sort(1, Qt.AscendingOrder)
        self.assertEqual([proxy.data(proxy.index(i, 0)) for i in range(4)], ["C", "B", "A", "D"])
        path = self.root / "comparison.csv"
        export_csv(path, proxy)
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertIn("Model!C4", rows[0][1])
        self.assertEqual(rows[1][1], "-2")
        proxy.setFilterFixedString("B")
        # Filter only the company column so a temporary path cannot match.
        proxy.setFilterKeyColumn(0)
        export_csv(path, proxy)
        with path.open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(len(list(csv.reader(stream))), 2)
        self.assertEqual(csv_value("=1+2"), "'=1+2")
        self.assertEqual(csv_value(-2), -2)

    def test_120_files_and_dialog_navigation(self):
        paths = []
        for i in range(120):
            path = self.root / f"Company_{i:03}.xlsx"
            make_workbook(path, value=i)
            paths.append(path)
        metrics = [self.metric, Metric("Margin", 4, 2)]
        start = time.perf_counter()
        dialog = ComparisonDialog([(p.stem, p) for p in paths], self.reference, "Model", metrics)
        dialog.show()
        wait_until(lambda: dialog.worker is None)
        self.assertEqual(dialog.model.rowCount(), 120)
        self.assertTrue(all(r.status == "OK" for r in dialog.model.rows))
        self.assertEqual(dialog.model.rows[-1].values, [119, 0.125])
        opened = []
        dialog.open_source.connect(lambda *args: opened.append(args))
        dialog.navigate(dialog.proxy.index(0, 1))
        self.assertEqual(opened[0][1:], ("Model", 3, 2))
        print(f"120 synthetic workbooks, 2 metrics: {time.perf_counter() - start:.3f}s")
        dialog.close()
        APP.sendPostedEvents()

    def test_cancel_and_close_during_read(self):
        dialog = ComparisonDialog([(str(i), self.reference) for i in range(500)],
                                  self.reference, "Model", [self.metric])
        dialog.show()
        dialog.cancel()
        wait_until(lambda: dialog.worker is None)
        self.assertIn("Cancelled", dialog.status.text())
        dialog.start()
        dialog.close()
        wait_until(lambda: dialog.worker is None)
        self.assertTrue(dialog._closing)

    def test_generic_folder_bounded_cache_and_source_navigation(self):
        from theia.app import MainWindow

        QSettings.setDefaultFormat(QSettings.IniFormat)
        QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(self.root / "settings"))
        for i in range(6):
            make_workbook(self.root / f"Company_{i:03}.xlsx", value=i)
        make_workbook(self.root / "~$temporary.xlsx")
        window = MainWindow()
        window.show()
        window.open_folder(self.root)
        try:
            companies = window._companies_model.findItems("Company*", Qt.MatchWildcard | Qt.MatchRecursive)
            self.assertEqual(len(companies), 6)
            for item in companies:
                window.company_tree.setCurrentIndex(window._companies_proxy.mapFromSource(item.index()))
            self.assertLessEqual(len(window._workbooks), 3)
            self.assertEqual(window.table.model().raw_value(3, 2), 5)
            index = window.table.model().index(3, 2)
            window.table.br.selectionModel().setCurrentIndex(index, QItemSelectionModel.ClearAndSelect)
            window._compare_selected()
            dialog = window._comparison_dialogs[-1]
            wait_until(lambda: dialog.worker is None)
            self.assertEqual(dialog.model.rowCount(), 6)
            window._open_comparison_source(companies[0].data(Qt.UserRole), "Model", 3, 2)
            self.assertEqual(window.table.currentIndex().row(), 3)
            self.assertEqual(window.table.currentIndex().column(), 2)
            window.company_filter.setText("Company_000")
            window._compare_selected()
            filtered = window._comparison_dialogs[-1]
            wait_until(lambda: filtered.worker is None)
            self.assertEqual(filtered.model.rowCount(), 1)
            filtered.close()
            window.open_files([self.reference, self.root / "Company_001.xlsx"])
            self.assertEqual(window._companies_model.rowCount(), 2)
            archive = self.root / "archive.zip"
            with zipfile.ZipFile(archive, "w") as z:
                z.write(self.reference, "Company_archive.xlsx")
            window.company_filter.clear()
            window.open_files([self.reference, archive])
            archive_items = window._companies_model.findItems("Company_archive", Qt.MatchRecursive | Qt.MatchExactly)
            self.assertEqual(len(archive_items), 1)
            window.company_tree.setCurrentIndex(window._companies_proxy.mapFromSource(archive_items[0].index()))
            self.assertEqual(window.table.model().raw_value(3, 2), 0)
            dialog.close()
        finally:
            window.close()


if __name__ == "__main__":
    unittest.main()
