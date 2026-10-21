import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook as ExcelWorkbook
from PySide6.QtCore import QCoreApplication, QEvent, QItemSelectionModel, QSettings, Qt
from PySide6.QtGui import QAction, QFontMetricsF
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication
from theia.app import MainWindow
from theia.workbook import Workbook

APP = QApplication.instance() or QApplication([])


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        settings = QSettings(str(self.root / "settings.ini"), QSettings.IniFormat)
        patcher = patch("theia.app.QSettings", return_value=settings)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.window = MainWindow()
        self.window.show()

    def tearDown(self):
        self.window.close()
        self.temporary.cleanup()

    def make_model(self, name):
        path = self.root / f"{name}.xlsx"
        workbook = ExcelWorkbook()
        sheet = workbook.active
        sheet.title = "Forecast"
        sheet.merge_cells("A1:D1")
        sheet["A1"] = "Generated company model"
        sheet["A2"] = "Revenue"
        sheet["B2"] = 125
        workbook.create_sheet("Valuation")["C3"] = 0
        workbook.save(path)
        workbook.close()
        return path

    def test_tabs_cell_navigation_and_merged_title_without_sidecars(self):
        self.assertEqual(self.window._sheet_stack.currentIndex(), 0)
        self.window.open_path(self.make_model("Company"))
        APP.processEvents()
        self.assertEqual(self.window._sheet_stack.currentIndex(), 1)
        self.assertEqual(self.window.sheet_tabs.tabText(0), "Forecast")
        self.assertEqual(self.window.table.freezeCols(), 0)
        self.assertEqual(self.window.table.tr.columnSpan(0, 0), 4)
        self.window._cell_address.setText("B2")
        self.window._go_to_cell()
        self.assertEqual(self.window._cell_value.text(), "125")
        self.window.sheet_tabs.setCurrentIndex(1)
        self.assertEqual(self.window._sheets_proxy.data(self.window.sheet_list.currentIndex()), "Valuation")
        self.window._cell_address.setText("C3")
        self.window._go_to_cell()
        self.assertEqual(self.window._cell_value.text(), "0")
        self.window.sheet_filter.setText("Forecast")
        APP.processEvents()
        self.assertEqual(self.window.sheet_tabs.count(), 1)
        self.assertEqual(self.window._sheets_proxy.data(self.window.sheet_list.currentIndex()), "Forecast")
        self.assertEqual(self.window.table.model().raw_value(0, 0), "Generated company model")
        self.window._set_sidebar_visible(False)
        self.assertFalse(self.window._nav_widgets[0].isVisible())
        self.assertTrue(self.window.sheet_tabs.isVisible())

    def test_opening_and_go_to_cell_return_keyboard_focus_to_the_grid(self):
        self.window.open_path(self.make_model("Company"))
        APP.processEvents()
        self.assertTrue(self.window.table.isAncestorOf(APP.focusWidget()))
        QTest.keyClick(APP.focusWidget(), Qt.Key_Down)
        QTest.keyClick(APP.focusWidget(), Qt.Key_Right)
        self.assertEqual(self.window._cell_address.text(), "B2")
        self.window._cell_address.setFocus()
        self.window._cell_address.setText("A1")
        QTest.keyClick(self.window._cell_address, Qt.Key_Return)
        self.assertTrue(self.window.table.isAncestorOf(APP.focusWidget()))
        QTest.keyClick(APP.focusWidget(), Qt.Key_Down)
        self.assertEqual(self.window._cell_address.text(), "A2")

    def test_clicking_a_sheet_tab_returns_keyboard_focus_to_the_grid(self):
        self.window.open_path(self.make_model("Company"))
        APP.processEvents()
        QTest.mouseClick(self.window.sheet_tabs, Qt.LeftButton, pos=self.window.sheet_tabs.tabRect(1).center())
        APP.processEvents()
        self.assertEqual(self.window.sheet_tabs.currentIndex(), 1)
        self.assertTrue(self.window.table.isAncestorOf(APP.focusWidget()))
        QTest.keyClick(APP.focusWidget(), Qt.Key_Down)
        self.assertEqual(self.window._cell_address.text(), "A2")

    def test_navigation_keeps_data_and_metadata_caches_bounded(self):
        for index in range(9):
            path = self.make_model(f"Company_{index}")
            path.with_suffix(".json").write_text(json.dumps({"sheets": {}}), encoding="utf-8")
        self.window.open_folder(self.root)
        companies = [
            item
            for item in self.window._companies_model.findItems("*", Qt.MatchWildcard | Qt.MatchRecursive)
            if item.data(Qt.UserRole) is not None
        ]
        for item in companies:
            self.window.company_tree.setCurrentIndex(self.window._companies_proxy.mapFromSource(item.index()))
            for row in range(self.window.sheet_tabs.count()):
                self.window.sheet_tabs.setCurrentIndex(row)
        self.assertLessEqual(len(self.window._workbooks), 1)
        self.assertLessEqual(len(self.window._sheet_model_cache), 4)
        self.assertLessEqual(len(self.window._sidecar_cache), 2)
        self.assertLessEqual(len(self.window._col_width_cache), 8)
        self.assertIsNone(self.window._preload_worker)
        QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
        self.assertEqual(len(self.window.table.findChildren(QItemSelectionModel)), 1)

    def test_long_history_columns_fit_displayed_numbers_and_keep_labels_readable(self):
        path = self.root / "Company_history.xlsx"
        workbook = ExcelWorkbook()
        sheet = workbook.active
        sheet.title = "History"
        sheet["A1"] = "Generated long-term financial statements"
        sheet["C1"] = "Values in millions; highlighted periods are forecasts"
        sheet["A3"], sheet["B3"] = "Metric", "Unit"
        sheet["A6"], sheet["B6"] = "販売費・一般管理費", "百万円"
        for column in range(3, 35):
            sheet.cell(3, column, f"{1988 + column}-03")
            sheet.cell(4, column, "Actual")
            sheet.cell(6, column, 12345 + column)
            sheet.cell(7, column, 0.12345678901234567)
            sheet.cell(8, column, -0.01234567890123456)
        workbook.save(path)
        workbook.close()
        self.window.open_path(path)
        APP.processEvents()
        table = self.window.table
        numeric_widths = [table.columnWidth(c) for c in range(2, 34)]
        self.assertTrue(all(60 <= width <= 100 for width in numeric_widths))
        metrics = QFontMetricsF(table.br.font())
        self.assertGreaterEqual(table.columnWidth(0), metrics.horizontalAdvance("販売費・一般管理費") + 6)
        for column in range(2, 34):
            for row in (2, 3, 5, 6, 7):
                text = table.model().data(table.model().index(row, column), Qt.DisplayRole)
                self.assertGreaterEqual(table.columnWidth(column), metrics.horizontalAdvance(text) + 6)
        self.assertAlmostEqual(table.model().raw_value(6, 2), 0.12345678901234567, places=15)

    def test_fit_columns_restores_content_widths_without_changing_selection_or_zoom(self):
        self.window.open_path(self.make_model("Company"))
        self.window._set_zoom(150)
        self.window._fit_columns()
        table = self.window.table
        expected = table.columnWidth(1)
        current, model = table.currentIndex(), table.model()
        table.setColumnWidth(1, 220)
        self.assertEqual(table.columnWidth(1), 330)
        action = next(a for a in self.window.findChildren(QAction) if a.text() == "Fit column &widths")
        action.trigger()
        self.assertEqual(table.columnWidth(1), expected)
        self.assertEqual(table.currentIndex(), current)
        self.assertIs(table.model(), model)
        self.assertEqual(self.window._zoom_percent, 150)

    def test_merged_titles_do_not_expand_individual_columns(self):
        self.window.open_path(self.make_model("Company"))
        widths = [self.window.table.columnWidth(c) for c in range(4)]
        self.assertTrue(all(width <= 100 for width in widths))

    def test_plain_workbooks_do_not_load_arrow_runtime(self):
        path = self.make_model("Plain")
        code = (
            "import sys; from PySide6.QtWidgets import QApplication; "
            "from theia.app import MainWindow; from pathlib import Path; "
            "app=QApplication([]); window=MainWindow(); window.open_path(Path(sys.argv[1])); "
            "assert 'pyarrow' not in sys.modules; window.close()"
        )
        subprocess.run([sys.executable, "-c", code, str(path)], check=True, capture_output=True)

    def make_formatted_model(self):
        path = self.root / "Company_formats.xlsx"
        workbook = ExcelWorkbook()
        sheet = workbook.active
        sheet.title = "Rates"
        sheet.append(["Measure", "Value"])
        for row, (label, value, fmt) in enumerate(
            [
                ("Planned profit growth", 0.125, "0.0%;[Red]-0.0%"),
                ("Margin", -0.25, "0%"),
                ("Growth", 0, "0.00%"),
                ("Value per share", 0.1234, "0.00"),
                ("Flag", True, "0%"),
                ("Text", "12%", "0%"),
                ("Uncalculated formula", "=1/8", "0.0%"),
            ],
            start=2,
        ):
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, value).number_format = fmt
        amounts = workbook.create_sheet("Amounts")
        amounts["A1"] = "Revenue"
        amounts["B1"] = 1234.5
        amounts["B1"].number_format = "#,##0.0"
        workbook.save(path)
        workbook.close()
        return path

    def test_source_number_formats_display_percentages_without_changing_values(self):
        self.window.open_path(self.make_formatted_model())
        model = self.window.table.model()
        for row, expected in enumerate(["12.5%", "-25%", "0.00%", "0.12", "TRUE", "12%", ""], start=1):
            self.assertEqual(model.data(model.index(row, 1), Qt.DisplayRole), expected)
        self.assertEqual(model.raw_value(1, 1), 0.125)
        self.assertEqual(model.raw_value(2, 1), -0.25)
        self.window.table.setCurrentIndex(model.index(1, 1))
        self.assertEqual(self.window._cell_value.text(), "12.5%")
        self.window.sheet_tabs.setCurrentIndex(1)
        model = self.window.table.model()
        self.assertEqual(model.data(model.index(0, 1), Qt.DisplayRole), "1,234.5")
        self.window.sheet_tabs.setCurrentIndex(0)
        model = self.window.table.model()
        self.assertEqual(model.data(model.index(1, 1), Qt.DisplayRole), "12.5%")

    def test_source_formats_are_read_from_workbooks_inside_zip_archives(self):
        path = self.make_formatted_model()
        archive = self.root / "Models.zip"
        with zipfile.ZipFile(archive, "w") as stream:
            stream.write(path, "models/Company.xlsx")
        self.window.open_path(archive)
        model = self.window.table.model()
        self.assertEqual(model.data(model.index(1, 1), Qt.DisplayRole), "12.5%")

    def test_explicit_sidecar_formats_override_workbook_formats(self):
        path = self.make_formatted_model()
        path.with_suffix(".json").write_text(
            json.dumps(
                {
                    "sheets": {
                        "Rates": {
                            "row_formats": {"2": "0.00%", "3": "0%"},
                            "cell_formats": {"3,2": "0.0%"},
                        }
                    }
                }
            ),
            encoding="utf-8",
        )
        self.window.open_path(path)
        model = self.window.table.model()
        self.assertEqual(model.data(model.index(1, 1), Qt.DisplayRole), "12.50%")
        self.assertEqual(model.data(model.index(2, 1), Qt.DisplayRole), "-25.0%")

    def test_source_formats_are_lazy_and_evicted_with_sheet_data(self):
        with Workbook(self.make_formatted_model()) as workbook:
            first = workbook.load_sheet("Rates")
            self.assertIsNone(workbook._number_formats)
            self.assertIsNone(workbook._xlsx)
            formats = workbook.cell_formats("Rates")
            self.assertEqual(formats[(2, 2)], "0.0%;[Red]-0.0%")
            self.assertIs(workbook.cell_formats("Rates"), formats)
            workbook.load_sheet("Amounts")
            self.assertIsNone(first.cell_formats)
            self.assertEqual(workbook.cell_formats("Amounts")[(1, 2)], "#,##0.0")

    def test_arrow_display_metadata_is_loaded_when_present(self):
        # Load Arrow only for the fixture that needs display sidecars.
        import pyarrow as pa  # noqa: PLC0415

        path = self.make_model("With_metadata")
        batch = pa.record_batch([pa.array(["Prepared value"])], names=["v_0"])
        batch = batch.replace_schema_metadata({b"n_rows": b"1", b"n_cols": b"1"})
        with pa.OSFile(str(path.with_suffix(".Forecast.arrow")), "wb") as stream:
            with pa.ipc.new_file(stream, batch.schema) as writer:
                writer.write(batch)
        self.window.open_path(path)
        self.assertEqual(self.window.table.model().raw_value(0, 0), "Prepared value")


if __name__ == "__main__":
    unittest.main()
